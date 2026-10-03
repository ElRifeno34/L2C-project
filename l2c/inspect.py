"""Inspect local dataset structure without exporting drawing text or cell values."""
import argparse
import json
import time
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


def inspect_dataset(root: Path, sample_text: bool = False) -> dict:
    projects = []
    for project in sorted(p for p in root.iterdir() if p.is_dir()):
        files = []
        for pdf in sorted(project.rglob("*.pdf")):
            entry = {"file": str(pdf.relative_to(root)), "bytes": pdf.stat().st_size}
            try:
                reader = PdfReader(pdf)
                entry["pages"] = len(reader.pages)
                entry["encrypted"] = reader.is_encrypted
                if sample_text and reader.pages:
                    text = reader.pages[0].extract_text() or ""
                    entry["first_page_text_characters"] = len(text.strip())
                    # A heuristic, not a guarantee that OCR is required.
                    entry["first_page_low_text"] = len(text.strip()) < 100
            except Exception as exc:
                entry["error_type"] = type(exc).__name__
            files.append(entry)
        projects.append({
            "project": project.name,
            "pdf_count": len(files),
            "page_count": sum(f.get("pages", 0) for f in files),
            "files": files,
        })
    spreadsheets = []
    for path in sorted(root.rglob("*.xlsx")):
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            spreadsheets.append({
                "file": str(path.relative_to(root)),
                "sheets": [{"name": s.title, "rows": s.max_row,
                            "columns": s.max_column} for s in workbook],
            })
        finally:
            workbook.close()
    return {"projects": projects, "spreadsheets": spreadsheets}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--sample-text", action="store_true",
                        help="Locally measure first-page text; never export its content.")
    args = parser.parse_args()
    root = args.data_dir.resolve()
    output = args.output.resolve()
    repo = Path(__file__).resolve().parents[1]
    if not root.is_dir():
        parser.error("Dataset directory does not exist.")
    if output.is_relative_to(repo):
        parser.error("Save inspection results outside the repository.")
    started = time.monotonic()
    report = inspect_dataset(root, args.sample_text)
    report["elapsed_seconds"] = round(time.monotonic() - started, 2)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    for project in report["projects"]:
        errors = sum("error_type" in f for f in project["files"])
        print(f"{project['project']}: {project['pdf_count']} PDFs, "
              f"{project['page_count']} pages, {errors} read errors", flush=True)
    print(f"Inspection completed in {report['elapsed_seconds']} seconds.")


if __name__ == "__main__":
    main()
